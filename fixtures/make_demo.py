"""Generate the controlled demo pair from PRD §30 / PDD §31, plus its ground truth.

Run: python -m fixtures.make_demo
Writes demo/proposal_v3.docx, demo/proposal_v4.docx, demo/ground_truth.json
"""

from __future__ import annotations

import json
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK

OUT = Path(__file__).resolve().parent / "demo"

# Version A then Version B. Sections are deliberately renamed and renumbered in B so
# the demo exercises section alignment, not just text diffing.
VERSION_A: list[tuple[str, list[str]]] = [
    ("", ["Acme Website Redesign Proposal", "Prepared for Acme Industries Limited",
          "Version 3 — 12 August 2026"]),
    ("1. Introduction", [
        "This proposal sets out the terms on which Northwind Digital Limited (the Supplier) "
        "will deliver a website redesign for Acme Industries Limited (the Client).",
    ]),
    ("2. Project Scope", [
        "The Supplier shall deliver a full redesign of the Client's public website, "
        "including information architecture, visual design and front-end implementation.",
        "The Supplier shall provide on-site training for up to twelve Client staff.",
        "The Supplier shall provide security monitoring.",
        "The Supplier shall migrate all existing content from the current platform.",
    ]),
    ("3. Pricing", [
        "The total project price is ₦5,000,000 exclusive of applicable taxes.",
        "The price includes design, development, testing and deployment.",
    ]),
    ("4. Delivery", [
        "The Supplier shall complete the project within 45 days of the commencement date.",
        "The commencement date is 5 October 2026.",
    ]),
    ("5. Support", [
        "The Supplier shall provide support and maintenance for 12 months following launch.",
        "Support requests shall be acknowledged within 8 hours.",
    ]),
    ("6. Payment Terms", [
        "The Client shall pay 50% of the total price upfront on signature of this proposal.",
        "The Client shall pay invoices within 30 days of receipt.",
        "Late payment shall attract interest at 2% per month.",
    ]),
    ("7. Responsibilities", [
        "The Client shall nominate a single point of contact for the duration of the project.",
        "The Client shall provide brand assets within five days of the commencement date.",
    ]),
    ("8. Confidentiality", [
        "Each party shall keep the other party's confidential information confidential for "
        "three years following termination of this agreement.",
    ]),
    ("9. Exclusivity", [
        "The Supplier shall not provide website redesign services to any direct competitor "
        "of the Client for twelve months following launch.",
    ]),
    ("10. Termination", [
        "Either party may terminate this agreement by giving 14 days written notice.",
        "The Client shall pay for all work completed up to the termination date.",
    ]),
    ("11. Governing Law", [
        "This agreement is governed by the laws of the Federal Republic of Nigeria.",
    ]),
]

VERSION_B: list[tuple[str, list[str]]] = [
    ("", ["Acme Website Redesign Proposal", "Prepared for Acme Industries Limited",
          "Version 4 — 2 September 2026"]),
    ("1. Introduction", [
        "This proposal sets out the terms on which Northwind Digital Limited (the Supplier) "
        "will deliver a website redesign for Acme Industries Limited (the Client).",
    ]),
    ("2. Scope of Services", [
        "The Supplier shall deliver a full redesign of the Client's public website, "
        "including information architecture, visual design and front-end implementation.",
        "The Supplier may provide security monitoring when requested.",
        "The Supplier shall migrate all existing content from the current platform.",
        "The Supplier shall provide monthly analytics reporting to the Client.",
    ]),
    ("3. Commercial Terms", [
        "The total project price is ₦6,000,000 exclusive of applicable taxes.",
        "The price includes design, development, testing and deployment.",
    ]),
    ("4. Timeline", [
        "The Supplier shall complete the project within 60 days of the commencement date.",
        "The commencement date is 12 October 2026.",
    ]),
    ("5. Support", [
        "The Supplier shall provide support and maintenance for 6 months following launch.",
        "Support requests shall be acknowledged within 8 hours.",
    ]),
    ("6. Payment", [
        "The Client shall pay 70% of the total price upfront on signature of this proposal.",
        "The Client shall pay invoices within 60 days of receipt.",
        "Late payment shall attract interest at 2% per month.",
    ]),
    ("7. Responsibilities", [
        "The Client shall nominate a single point of contact for the duration of the project.",
        "The Client shall provide brand assets within five days of the commencement date.",
        "The Client must provide weekly access reports to the Supplier.",
    ]),
    ("8. Confidentiality", [
        "Each party shall keep the other party's confidential information confidential for "
        "three years following termination of this agreement.",
    ]),
    ("9. Termination", [
        "Either party may terminate this agreement by giving 30 days written notice.",
        "The Client shall pay for all work completed up to the termination date.",
    ]),
    ("10. Governing Law", [
        "This agreement is governed by the laws of the Federal Republic of Nigeria.",
    ]),
    ("11. Data Protection", [
        "The Supplier shall process personal data only on the documented instructions of "
        "the Client and shall delete all personal data within 30 days of termination.",
    ]),
]

# Rate-card table, to exercise structural table diffing (TRD §32).
TABLE_A = [
    ["Service", "Price"],
    ["Consulting", "$10,000"],
    ["Training", "$5,000"],
    ["Deployment", "$3,000"],
]
TABLE_B = [
    ["Service", "Price"],
    ["Consulting", "$12,500"],
    ["Deployment", "$3,000"],
    ["Support", "$4,000"],
]

GROUND_TRUTH = [
    {"id": "TEST-001", "category": "PRICING", "importance": "HIGH",
     "expect": {"old": 5000000, "new": 6000000, "absolute": 1000000, "percentage": 20.0},
     "match": ["5,000,000", "6,000,000"],
     "note": "Total project price ₦5,000,000 -> ₦6,000,000"},
    {"id": "TEST-002", "category": "DURATION", "importance": "HIGH",
     "expect": {"old": 45, "new": 60, "unit": "days"},
     "match": ["45 days", "60 days"],
     "note": "Delivery period 45 -> 60 days"},
    {"id": "TEST-003", "category": "SCOPE", "importance": "HIGH",
     "expect": {"removed": "on-site training"},
     "match": ["on-site training"],
     "note": "On-site training clause removed"},
    {"id": "TEST-004", "category": "DURATION", "importance": "HIGH",
     "expect": {"old": 12, "new": 6, "unit": "months"},
     "match": ["12 months", "6 months"],
     "note": "Support period 12 -> 6 months"},
    {"id": "TEST-005", "category": "PAYMENT_TERMS", "importance": "HIGH",
     "expect": {"old": 50, "new": 70, "unit": "percent"},
     "match": ["50%", "70%"],
     "note": "Upfront payment 50% -> 70%"},
    {"id": "TEST-006", "category": "TERMINATION", "importance": "HIGH",
     "expect": {"old": 14, "new": 30, "unit": "days"},
     "match": ["14 days", "30 days"],
     "note": "Termination notice 14 -> 30 days"},
    {"id": "TEST-007", "category": "PAYMENT_TERMS", "importance": "HIGH",
     "expect": {"old": 30, "new": 60, "unit": "days"},
     "match": ["within 30 days", "within 60 days"],
     "note": "Invoice payment period 30 -> 60 days"},
    {"id": "TEST-008", "category": "OBLIGATIONS", "importance": "HIGH",
     "expect": {"modality_old": "shall", "modality_new": "may"},
     "match": ["security monitoring"],
     "note": "Security monitoring: shall -> may when requested (meaning change)"},
    {"id": "TEST-009", "category": "SCOPE", "importance": "MEDIUM",
     "expect": {"added": "monthly analytics reporting"},
     "match": ["analytics reporting"],
     "note": "Monthly analytics reporting added"},
    {"id": "TEST-010", "category": "RESPONSIBILITIES", "importance": "MEDIUM",
     "expect": {"added": "weekly access reports"},
     "match": ["weekly access reports"],
     "note": "New Client obligation: weekly access reports"},
    {"id": "TEST-011", "category": "SCOPE", "importance": "HIGH",
     "expect": {"removed_section": "Exclusivity"},
     "match": ["Exclusivity", "direct competitor"],
     "note": "Exclusivity section removed"},
    {"id": "TEST-012", "category": "CONFIDENTIALITY", "importance": "MEDIUM",
     "expect": {"added_section": "Data Protection"},
     "match": ["Data Protection", "personal data"],
     "note": "Data Protection section added"},
    {"id": "TEST-017", "category": "DATES", "importance": "LOW",
     "expect": {"old": "2026-08-12", "new": "2026-09-02", "days": 21},
     "match": ["12 August 2026", "2 September 2026"],
     "note": "Cover page version date 12 August -> 2 September 2026"},
    {"id": "TEST-013", "category": "DATES", "importance": "MEDIUM",
     "expect": {"old": "2026-10-05", "new": "2026-10-12", "days": 7},
     "match": ["5 October 2026", "12 October 2026"],
     "note": "Commencement date moved 7 days later"},
    {"id": "TEST-014", "category": "PRICING", "importance": "MEDIUM",
     "expect": {"old": 10000, "new": 12500},
     "match": ["10,000", "12,500"],
     "note": "Rate card: Consulting $10,000 -> $12,500"},
    {"id": "TEST-015", "category": "PRICING", "importance": "MEDIUM",
     "expect": {"removed_row": "Training"},
     "match": ["Training"],
     "note": "Rate card: Training row removed"},
    {"id": "TEST-016", "category": "PRICING", "importance": "MEDIUM",
     "expect": {"added_row": "Support"},
     "match": ["Support"],
     "note": "Rate card: Support row added"},
]


def build(sections: list[tuple[str, list[str]]], table: list[list[str]], path: Path) -> None:
    doc = docx.Document()
    for index, (heading, paragraphs) in enumerate(sections):
        if heading:
            doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)
        # Rate card sits inside the pricing section.
        if heading.endswith(("Pricing", "Commercial Terms")):
            doc.add_paragraph("Rate card:")
            t = doc.add_table(rows=len(table), cols=len(table[0]))
            t.style = "Table Grid"
            for r, row in enumerate(table):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
        # Force a page break part-way so page citations are meaningful.
        if index == len(sections) // 2:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main() -> None:
    build(VERSION_A, TABLE_A, OUT / "proposal_v3.docx")
    build(VERSION_B, TABLE_B, OUT / "proposal_v4.docx")
    (OUT / "ground_truth.json").write_text(
        json.dumps(
            {
                "pair": "demo-proposal",
                "version_a": "proposal_v3.docx",
                "version_b": "proposal_v4.docx",
                "changes": GROUND_TRUTH,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"wrote {OUT / 'proposal_v3.docx'}")
    print(f"wrote {OUT / 'proposal_v4.docx'}")
    print(f"wrote {OUT / 'ground_truth.json'} ({len(GROUND_TRUTH)} labelled changes)")


if __name__ == "__main__":
    main()
